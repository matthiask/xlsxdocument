import datetime as dt
import io
from decimal import Decimal

from django.contrib import admin
from django.contrib.auth.models import Group
from django.test import TestCase
from openpyxl import load_workbook

from xlsxdocument import XLSXDocument, create_export_selected, export_selected


def read(response):
    workbook = load_workbook(io.BytesIO(response.content))
    return {
        sheet.title: [[cell.value for cell in row] for row in sheet.iter_rows()]
        for sheet in workbook.worksheets
    }


class XLSXDocumentTest(TestCase):
    def test_table(self):
        xlsx = XLSXDocument()
        xlsx.add_sheet("Sales")
        xlsx.table(["Product", "Quantity", "Price"], [["Widget", 10, 9.99]])
        response = xlsx.to_response("sales.xlsx")

        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="sales.xlsx"'
        )
        self.assertEqual(
            read(response),
            {"Sales": [["Product", "Quantity", "Price"], ["Widget", 10, 9.99]]},
        )

    def test_multiple_sheets(self):
        xlsx = XLSXDocument()
        xlsx.add_sheet("One")
        xlsx.table([], [["a"]])
        xlsx.add_sheet("Two")
        xlsx.table([], [["b"]])
        self.assertEqual(
            read(xlsx.to_response("export.xlsx")), {"One": [["a"]], "Two": [["b"]]}
        )

    def test_values(self):
        xlsx = XLSXDocument()
        xlsx.add_sheet("Values")
        xlsx.table(
            [],
            [
                [
                    None,
                    Decimal("1.5"),
                    dt.date(2026, 10, 28),
                    dt.datetime(2026, 10, 28, 12, tzinfo=dt.timezone.utc),
                    dt.datetime(2026, 10, 28, 12),
                    "  padded\x01\x0b ",
                    ["a", "b"],
                ]
            ],
        )
        self.assertEqual(
            read(xlsx.to_response("values.xlsx")),
            {
                "Values": [
                    [
                        "-",
                        1.5,
                        dt.datetime(2026, 10, 28),
                        # Aware datetimes are converted to the current time zone
                        dt.datetime(2026, 10, 28, 13),
                        dt.datetime(2026, 10, 28, 12),
                        "padded",
                        "['a', 'b']",
                    ]
                ]
            },
        )


class ExportSelectedTest(TestCase):
    modeladmin = admin.ModelAdmin(Group, admin.site)

    def test_export_selected(self):
        Group.objects.create(name="Editors")
        response = export_selected(self.modeladmin, None, Group.objects.all())
        self.assertEqual(
            response["Content-Disposition"], 'attachment; filename="auth.group.xlsx"'
        )
        rows = read(response)["groups"]
        self.assertEqual(rows[0], ["__str__", "ID", "Name"])
        self.assertEqual(rows[1][0::2], ["Editors", "Editors"])

    def test_create_export_selected(self):
        Group.objects.create(name="Editors")
        action = create_export_selected(
            additional=[("Length", lambda group: len(group.name))]
        )
        response = action(self.modeladmin, None, Group.objects.all())
        rows = read(response)["groups"]
        self.assertEqual(rows[0], ["__str__", "ID", "Name", "Length"])
        self.assertEqual(rows[1][3], 7)

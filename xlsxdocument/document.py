import datetime as dt
import io
import re
from decimal import Decimal

from django.http import HttpResponse
from django.utils.text import capfirst, slugify
from django.utils.timezone import is_naive, make_naive
from django.utils.translation import gettext_lazy as _
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell


__all__ = ("Formula", "XLSXDocument", "create_export_selected", "export_selected")


# openpyxl, or maybe XLSX doesn't like all characters
ILLEGAL_CHARACTERS_RE = re.compile(r"[\000-\010]|[\013-\014]|[\016-\037]")


class Formula(str):
    """Marks a string such as ``"=SUM(A1:A3)"`` as an intentional formula

    Other strings starting with ``=`` are written as text since they may
    contain untrusted data (formula injection).
    """


class XLSXDocument:
    def __init__(self):
        self.workbook = Workbook(write_only=True)
        self.sheet = None

    def add_sheet(self, title=None):
        self.sheet = self.workbook.create_sheet(title=title)

    def table(self, titles, rows):
        if titles:
            self.sheet.append(self._process(titles))

        for row in rows:
            self.sheet.append(self._process(row))

    def _process(self, row):
        processed = []
        for value in row:
            if isinstance(value, dt.datetime):
                processed.append(value if is_naive(value) else make_naive(value))
            elif isinstance(value, (int, float, Decimal, dt.date, Formula)):
                processed.append(value)
            elif value is None:
                processed.append("-")
            else:
                value = ILLEGAL_CHARACTERS_RE.sub("", ("%s" % value).strip())
                if value.startswith("="):
                    # openpyxl would write the value as a formula
                    cell = WriteOnlyCell(self.sheet, value)
                    cell.data_type = "s"
                    cell.quotePrefix = True
                    value = cell
                processed.append(value)
        return processed

    def table_from_queryset(self, queryset, additional=()):
        opts = queryset.model._meta

        titles = ["__str__"]
        titles.extend(str(capfirst(field.verbose_name)) for field in opts.fields)
        titles.extend(a[0] for a in additional)

        data = []
        for instance in queryset:
            row = ["%s" % instance]
            for field in opts.fields:
                if field.choices:
                    row.append(getattr(instance, "get_%s_display" % field.name)())
                else:
                    row.append(getattr(instance, field.name))

            row.extend(a[1](instance) for a in additional)
            data.append(row)

        self.add_sheet(slugify("%s" % opts.verbose_name_plural))
        self.table(titles, data)

    def to_response(self, filename):
        with io.BytesIO() as buf:
            self.workbook.save(buf)
            response = HttpResponse(
                buf.getvalue(),
                content_type=(
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                ),
            )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


_description = _("export_selected")


def create_export_selected(additional=(), short_description=_description):
    def export_selected(modeladmin, request, queryset):
        xlsx = XLSXDocument()
        xlsx.table_from_queryset(queryset, additional=additional)
        return xlsx.to_response(
            "%s.%s.xlsx"
            % (modeladmin.model._meta.app_label, modeladmin.model._meta.model_name)
        )

    export_selected.short_description = short_description

    return export_selected


export_selected = create_export_selected()
